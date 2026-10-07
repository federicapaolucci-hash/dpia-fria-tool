"""Export the frozen T17 workbook to config/*.yaml, with cross-reference validation.

    python tools/export_workbook.py                 # validate, then write config/
    python tools/export_workbook.py --check         # validate, compare with config/ on disk, write nothing
    python tools/export_workbook.py --encoding-report

The workbook is the only source of logic. What it leaves in prose (triggers, visibility,
option lists) is hand-coded in tools/encodings/*.yaml and merged here; every choice cites
an AS-* entry of spec/ANOMALIE_REV2.md. Unknown labels, unknown IDs and unexpected sheet
layouts are errors: nothing is guessed. openpyxl is used only here, never at runtime.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from pydantic import ValidationError

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from digcon.domain.enums import (  # noqa: E402
    HumanGateRequirement,
    MitigationAllowed,
    MitigationHierarchy,
    MitigationRequirement,
    MitigationStatus,
    MitigationTiming,
    DecisionEffect,
    Phase,
    Probability,
    Reversibility,
    RiskLevel,
    RuleSeverity,
    Scale,
    Scope,
    Severity,
    parse,
)
from digcon.domain.models import ConfigBundle  # noqa: E402
from digcon.domain.validation import Finding, validate_bundle  # noqa: E402

WORKBOOK_VERSION = "T17-REV2"
DEFAULT_WORKBOOK = ROOT.parent / "spec" / "DIGCON_T17_DECISION_ENGINE_REV_2.xlsx"
DEFAULT_ANOMALIES = ROOT.parent / "spec" / "ANOMALIE_REV2.md"
ENCODINGS = ROOT / "tools" / "encodings"
DEFAULT_OUT = ROOT / "config"

QUESTION_HEADERS = [
    "Fase", "Sottosezione", "ID", "Domanda / campo", "Modello risposta", "Visibilità / dipendenza",
    "Se positivo / adeguato", "Se negativo / criticità", "Se unknown / N/A / not owned",
    "Rule ID / dipendenza", "Stato revisione RA",
]  # fmt: skip
RULE_HEADERS = [
    "Rule ID", "Fonte", "Question / Risk link", "Attore", "Trigger / elementi", "Missing/unknown treatment",
    "Consequence class", "Severity", "Effetto sul flusso", "Può produrre O5?", "Human gate", "Stato / nota RA",
    "SEP implication", "Mitigation allowed?", "Mitigation output / action", "Remediability", "SEP / Mitigation ID",
]  # fmt: skip
RISK_HEADERS_USED = {
    0: "Risk ID", 1: "Fonte domanda/e", 2: "Attore / contesto", 3: "Diritto / area", 4: "Persone / gruppi",
    5: "Harm scenario", 10: "Relevant safeguards / controls to verify in the case",
    12: "Candidate mitigation / safeguard (pre-mapped)", 16: "Rule / HS link", 17: "Human review? [rule-driven]",
    18: "Template / case status", 19: "SEP relevant?", 23: "Mitigation ID",
    29: "RA reference profile (legacy — non-engine)", 30: "RA provisional residual risk (legacy — non-engine)",
    35: "PROP rule(s)",
}  # fmt: skip
MITIGATION_HEADERS = [
    "Mitigation ID", "Source type", "Source ID(s)", "Risk ID", "Source rule / HS ID (PROP-* / RISK-* / other rule)",
    "SEP ID", "Harm / issue addressed", "Fundamental right(s)", "Trigger / reason", "Measure", "Mitigation hierarchy",
    "Measure type", "Responsible role", "Owner", "Timing", "Requirement status", "Decision effect",
    "Evidence of implementation", "Verification method", "Closure criterion", "Status", "Risk before",
    "Expected residual risk", "Actual residual risk", "Mitigation adequacy", "Residual risk acceptable?",
    "Review trigger", "Review date", "Notes", "Expected effect", "Reassessment result", "Human legal review?",
    "Unresolved issue",
]  # fmt: skip
SEP_HEADERS = [
    "ID", "Field", "Question / information", "Response model", "Activation / dependency",
    "Output / decision effect", "Risk link", "Rule link", "Mitigation link", "Human review?", "Status",
]  # fmt: skip

PHASES = {
    "INTRO": Phase.INTRO, "WHAT": Phase.WHAT, "HOW": Phase.HOW, "WHY": Phase.WHY,
    "PROVIDER ADD-ON": Phase.PROVIDER, "DEPLOYER FRIA ADD-ON": Phase.DEPLOYER,
}  # fmt: skip

SEVERITY = {
    "0 — NONE": RuleSeverity.NONE, "1 — LOW": RuleSeverity.LOW, "2 — MEDIUM": RuleSeverity.MEDIUM,
    "3 — HIGH": RuleSeverity.HIGH, "4 — HARD STOP": RuleSeverity.HARD_STOP,
    "By RISK-*": RuleSeverity.BY_RISK, "By source RISK-*": RuleSeverity.BY_RISK,
    "By source rule": RuleSeverity.BY_SOURCE_RULE,
}  # fmt: skip
SEVERITY_NA_RE = re.compile(r"^N/A(?:\s+—\s+.+)?$")

# AS-005: unconditional YES -> REQUIRED, conditional formulations -> POSSIBLE, HS -> REQUIRED.
HUMAN_GATE = {
    "No": HumanGateRequirement.NO,
    "NO": HumanGateRequirement.NO,
    "Possible": HumanGateRequirement.POSSIBLE,
    "Possibile": HumanGateRequirement.POSSIBLE,
    "Possible if legal classification is uncertain": HumanGateRequirement.POSSIBLE,
    "Possible when legal characterisation of an element is contested": HumanGateRequirement.POSSIBLE,
    "Possible where legal reasoning requires validation": HumanGateRequirement.POSSIBLE,
    "Possible / YES for legal balancing where needed": HumanGateRequirement.POSSIBLE,
    "YES when normative/legal balancing is required": HumanGateRequirement.POSSIBLE,
    "YES where legal adequacy is contested": HumanGateRequirement.POSSIBLE,
    "YES where legal discrimination assessment is required": HumanGateRequirement.POSSIBLE,
    "YES where essentiality/non-remediability is a legal judgment": HumanGateRequirement.POSSIBLE,
    "YES for normative balancing where required": HumanGateRequirement.POSSIBLE,
    "YES where purpose legitimacy or normative suitability requires legal judgment": HumanGateRequirement.POSSIBLE,
    "YES when equivalence/suitability of alternatives requires normative assessment": HumanGateRequirement.POSSIBLE,
    "YES when output is HUMAN LEGAL REVIEW": HumanGateRequirement.POSSIBLE,
    "Sì per HS giuridici": HumanGateRequirement.POSSIBLE,
    "YES for legal HS": HumanGateRequirement.POSSIBLE,
    "As required by HS": HumanGateRequirement.POSSIBLE,
    "YES": HumanGateRequirement.REQUIRED,
    "Sì": HumanGateRequirement.REQUIRED,
    "Sì — legal gate": HumanGateRequirement.REQUIRED,
    "YES — residual-risk acceptability determination": HumanGateRequirement.REQUIRED,
}
HS_HUMAN_GATE = "Sì se la regola richiede valutazione giuridica"

MITIGATION_ALLOWED = {
    "YES": MitigationAllowed.YES, "NO": MitigationAllowed.NO, "CONDITIONAL": MitigationAllowed.CONDITIONAL,
    "N/A": MitigationAllowed.NOT_APPLICABLE, "YES BEFORE HS / NO ONCE MET": MitigationAllowed.BEFORE_HS_ONLY,
}  # fmt: skip
TIMING = {
    "Before pilot/deployment if activated": MitigationTiming.BEFORE_DEPLOYMENT,
    "Before pilot where material; repeat on review trigger": MitigationTiming.BEFORE_PILOT_AND_ON_REVIEW,
    "During pilot/deployment + continuous review": MitigationTiming.DURING_DEPLOYMENT_CONTINUOUS,
}
REQUIREMENT = {"Conditional": MitigationRequirement.CONDITIONAL, "Human determination": MitigationRequirement.HUMAN_DETERMINATION}
DECISION_EFFECT = {
    "Continuous condition if approved": DecisionEffect.CONTINUOUS_CONDITION,
    "Required before pilot/deployment if activated by rule/governance": DecisionEffect.REQUIRED_BEFORE_DEPLOYMENT,
    "Human governance decision required before activation": DecisionEffect.HUMAN_DECISION_BEFORE_ACTIVATION,
    "None until evidence gap is resolved": DecisionEffect.NONE_UNTIL_EVIDENCE_GAP_RESOLVED,
}
SEP_HUMAN_REVIEW = {"Possible": HumanGateRequirement.POSSIBLE, "No": HumanGateRequirement.NO}


class ExportError(Exception):
    pass


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


def text(value: Any) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def lookup(table: dict[str, Any], value: Any, where: str) -> Any:
    key = text(value)
    if key not in table:
        raise ExportError(f"{where}: unexpected value {key!r}; known: {sorted(table)}")
    return table[key]


_REF_RE = re.compile(
    r"(?P<range>\b(?P<rp>C|P|D|HS|G)(?P<ra>\d{2})\s*[-–]\s*(?:(?P=rp))?(?P<rb>\d{2})\b)"
    r"|(?P<wild>\b(?:RISK|MIT|SEP|PROP|EG|FRIA|RMS)-\*|\bHS\*)"
    r"|(?P<id>\bRISK-\d{2}\b|\bMIT-\d{3}\b|\bHS0[1-9]\b|\bRMS-9[A-Z]\b|\bFRIA-27[A-Z]\b|\b[A-Z]{2,5}-\d{2}\b"
    r"|\bC\d{2}B?\b|\b[NWQPDG]\d{2}\b|\bFOLLOW\b)"
)
_PROSE_GROUPS = {"RMS (tutta RMS)": "RMS-*", "(tutta FRIA)": "FRIA-*"}


def parse_refs(raw: Any) -> list[str]:
    """IDs cited in a workbook cell, ranges expanded, order kept, duplicates dropped."""
    s = text(raw)
    if s is None or s.upper() == "NESSUNA":
        return []
    for prose, token in _PROSE_GROUPS.items():
        s = s.replace(prose, f" {token} ")
    out: list[str] = []
    for m in _REF_RE.finditer(s):
        if m.group("range"):
            prefix, a, b = m.group("rp"), int(m.group("ra")), int(m.group("rb"))
            if b < a:
                raise ExportError(f"bad range {m.group('range')!r} in {raw!r}")
            out += [f"{prefix}{i:02d}" for i in range(a, b + 1)]
        else:
            out.append(m.group("wild") or m.group("id"))
    return list(dict.fromkeys(out))


def find_header(ws, expected: list[str] | dict[int, str]) -> int:
    want = expected if isinstance(expected, dict) else dict(enumerate(expected))
    for row in ws.iter_rows(min_row=1, max_row=40):
        if all(text(row[i].value) == h for i, h in want.items() if i < len(row)):
            return row[0].row
    raise ExportError(f"{ws.title}: header {list(want.values())[:3]}… not found — workbook layout changed?")


def data_rows(ws, header_row: int, id_col: int, id_re: str):
    pattern = re.compile(id_re)
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rid = text(row[id_col]) if len(row) > id_col else None
        if rid and pattern.match(rid):
            yield row


class StrictLoader(yaml.SafeLoader):
    """SafeLoader that refuses duplicate keys and reads YES/NO/ON/OFF as strings.

    YAML 1.1 turns a bare ``YES`` into ``True``: here only ``true``/``false`` are booleans,
    so answer states can be written unquoted. Merge keys (``<<``) may still override.
    """


StrictLoader.yaml_implicit_resolvers = {
    k: [(tag, rx) for tag, rx in v if tag != "tag:yaml.org,2002:bool"]
    for k, v in yaml.SafeLoader.yaml_implicit_resolvers.items()
}
StrictLoader.add_implicit_resolver(
    "tag:yaml.org,2002:bool", re.compile(r"^(?:true|True|TRUE|false|False|FALSE)$"), list("tTfF")
)


def _strict_mapping(loader: StrictLoader, node: yaml.MappingNode, deep: bool = False):
    seen: set[Any] = set()
    for key_node, _ in node.value:
        if key_node.tag == "tag:yaml.org,2002:merge":
            continue
        key = loader.construct_object(key_node, deep=True)
        if key in seen:
            raise yaml.constructor.ConstructorError(None, None, f"duplicate key {key!r}", key_node.start_mark)
        seen.add(key)
    return yaml.SafeLoader.construct_mapping(loader, node, deep=deep)


StrictLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, _strict_mapping)


def load_encoding(name: str) -> Any:
    with open(ENCODINGS / name, encoding="utf-8") as fh:
        return yaml.load(fh, Loader=StrictLoader)


# ---------------------------------------------------------------------------
# sheet readers
# ---------------------------------------------------------------------------


def read_questions(wb, enc: dict[str, Any], option_sets: dict[str, Any]) -> list[dict[str, Any]]:
    models = enc["response_models"]
    overrides: dict[str, dict[str, Any]] = enc.get("questions") or {}
    questions: list[dict[str, Any]] = []
    for sheet in ("01_CORE_COMUNE", "02_PROVIDER_ART9", "03_DEPLOYER_FRIA"):
        ws = wb[sheet]
        hdr = find_header(ws, QUESTION_HEADERS)
        for row in data_rows(ws, hdr, 2, r"^[A-Z]"):
            qid = text(row[2])
            model = text(row[4])
            base = lookup(models, model, f"{sheet}/{qid} response model")
            ov = overrides.get(qid, {})
            unknown_keys = set(ov) - {"kind", "states", "option_set", "visibility", "assumptions"}
            if unknown_keys:
                raise ExportError(f"encodings/questions.yaml {qid}: unknown keys {sorted(unknown_keys)}")
            questions.append(
                {
                    "id": qid,
                    "sheet": sheet,
                    "order": len(questions) + 1,
                    "phase": lookup(PHASES, row[0], f"{sheet}/{qid} phase"),
                    "subsection": text(row[1]),
                    "text": text(row[3]),
                    "response_model": model,
                    "kind": ov.get("kind", base["kind"]),
                    "states": ov.get("states", base["states"]),
                    "option_set": ov.get("option_set"),
                    "visibility_text": text(row[5]),
                    "visibility": ov.get("visibility"),
                    "if_positive": text(row[6]) or "",
                    "if_negative": text(row[7]) or "",
                    "if_unresolved": text(row[8]) or "",
                    "rule_refs_text": text(row[9]),
                    "rule_refs": parse_refs(row[9]),
                    "ra_status": text(row[10]),
                    "assumptions": ov.get("assumptions", []),
                }
            )
    known = {q["id"] for q in questions}
    stray = set(overrides) - known
    if stray:
        raise ExportError(f"encodings/questions.yaml: overrides for unknown questions {sorted(stray)}")
    for q in questions:
        if q["option_set"] and q["option_set"] not in option_sets:
            raise ExportError(f"{q['id']}: option set {q['option_set']} not in encodings/options.yaml")
    return questions


def read_governance(wb) -> list[dict[str, Any]]:
    ws = wb["00_MAPPA_FLUSSO"]
    hdr = find_header(ws, ["ID", "Field", "Purpose"])
    return [
        {"id": text(r[0]), "field": text(r[1]), "purpose": text(r[2])} for r in data_rows(ws, hdr, 0, r"^G\d{2}$")
    ]


def rule_family(rule_id: str) -> str:
    if re.match(r"^HS0[1-9]$", rule_id):
        return "HS"
    return rule_id.split("-")[0]


def rule_severity(raw: Any, where: str) -> RuleSeverity:
    s = text(raw)
    if s and SEVERITY_NA_RE.match(s):
        return RuleSeverity.NOT_APPLICABLE
    return lookup(SEVERITY, s, where)


def read_rules(wb, enc: dict[str, Any]) -> list[dict[str, Any]]:
    ws = wb["05_REGOLE_HS"]
    hdr = find_header(ws, RULE_HEADERS)
    encodings: dict[str, Any] = enc["rules"]
    rules: list[dict[str, Any]] = []
    for row in data_rows(ws, hdr, 0, r"^(?:HS0\d|[A-Z]+-[0-9A-Z]+)$"):
        rid = text(row[0])
        family = rule_family(rid)
        if rid not in encodings:
            raise ExportError(f"encodings/rules.yaml: no encoding for {rid}")
        encoding = dict(encodings[rid])
        gate_text = text(row[10])
        if gate_text == HS_HUMAN_GATE:
            if family != "HS":
                raise ExportError(f"{rid}: HS human-gate wording outside HS01–HS09")
            gate = HumanGateRequirement.REQUIRED  # AS-005
        else:
            gate = lookup(HUMAN_GATE, gate_text, f"{rid} human gate")
        o5_text = text(row[9])
        if o5_text == "NO":
            can_o5 = False
        elif o5_text and o5_text.startswith("YES"):
            can_o5 = True
        else:
            raise ExportError(f"{rid}: unexpected 'Può produrre O5?' {o5_text!r}")
        if family == "HS":
            encoding["elements"] = hs_elements(rid, text(row[4]), encoding.get("elements", []))
        rules.append(
            {
                "id": rid,
                "family": family,
                "order": len(rules) + 1,
                "source": text(row[1]),
                "links_text": text(row[2]),
                "link_ids": parse_refs(row[2]),
                "actor": text(row[3]),
                "trigger_text": text(row[4]),
                "missing_treatment": text(row[5]),
                "consequence_class": text(row[6]),
                "severity": rule_severity(row[7], f"{rid} severity"),
                "severity_text": text(row[7]),
                "flow_effect": text(row[8]),
                "can_produce_o5": can_o5,
                "human_gate": gate,
                "human_gate_text": gate_text,
                "ra_note": text(row[11]),
                "sep_implication": text(row[12]),
                "mitigation_allowed": lookup(MITIGATION_ALLOWED, row[13], f"{rid} mitigation allowed"),
                "mitigation_allowed_text": text(row[13]),
                "mitigation_action": text(row[14]),
                "remediability": text(row[15]),
                "related_text": text(row[16]),
                "related_ids": parse_refs(row[16]),
                "encoding": encoding,
            }
        )
    stray = set(encodings) - {r["id"] for r in rules}
    if stray:
        raise ExportError(f"encodings/rules.yaml: encodings for unknown rules {sorted(stray)}")
    return rules


def hs_elements(rid: str, trigger: str, encoded: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Constituent elements of a hard stop: text from the workbook, attributes from the encoding."""
    parts = [p.strip() for p in re.split(r"\s+AND\s+", trigger) if p.strip()]
    if len(parts) != len(encoded):
        raise ExportError(f"{rid}: workbook has {len(parts)} elements ({parts}), encoding has {len(encoded)}")
    out = []
    for i, (part, el) in enumerate(zip(parts, encoded), 1):
        if el.get("id") != f"{rid}.E{i}":
            raise ExportError(f"{rid}: element {i} must be {rid}.E{i}, got {el.get('id')}")
        if "text" in el:
            raise ExportError(f"{el['id']}: HS element text comes from the workbook, remove it from the encoding")
        out.append({**el, "text": part})
    return out


_LEGACY_RE = re.compile(
    r"scale=(?P<scale>[A-Z ]+); scope=(?P<scope>[A-Z ]+); reversibility=(?P<rev>[A-Z ]+); "
    r"probability=(?P<prob>[A-Z ]+); resulting reference risk=(?P<risk>[A-Z ]+)\."
)
_SOURCE_SECTION_RE = re.compile(r"^(Core|Provider|Deployer):\s*(.*)$")


def read_risks(wb) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    ws = wb["04_RISK_ANALYSIS"]
    hdr = find_header(ws, RISK_HEADERS_USED)
    risks = []
    for row in data_rows(ws, hdr, 0, r"^RISK-\d{2}\b"):
        cell = text(row[0])
        m = re.match(r"^(RISK-\d{2})(?:\s*\((.+)\))?$", cell)
        if not m:
            raise ExportError(f"04: cannot parse risk id {cell!r}")
        rid, label_ = m.group(1), m.group(2)
        sections: dict[str, list[str]] = {"Core": [], "Provider": [], "Deployer": []}
        for line in text(row[1]).splitlines():
            sm = _SOURCE_SECTION_RE.match(line.strip())
            if not sm:
                raise ExportError(f"{rid}: unexpected source line {line!r}")
            sections[sm.group(1)] = parse_refs(sm.group(2))
        legacy_text = text(row[29])
        lm = _LEGACY_RE.search(legacy_text or "")
        if not lm:
            raise ExportError(f"{rid}: cannot parse RA reference profile {legacy_text!r}")
        risks.append(
            {
                "id": rid,
                "label": label_,
                "order": len(risks) + 1,
                "source_text": text(row[1]),
                "core_questions": sections["Core"],
                "provider_questions": sections["Provider"],
                "deployer_questions": sections["Deployer"],
                "actor": text(row[2]),
                "right": text(row[3]),
                "groups": text(row[4]),
                "harm_scenarios": [ln.strip() for ln in text(row[5]).splitlines() if ln.strip()],
                "safeguards_to_verify": text(row[10]),
                "candidate_mitigation_text": text(row[12]),
                "rule_links_text": text(row[16]),
                "rule_links": parse_refs(row[16]),
                "human_review_text": text(row[17]),
                "template_status": text(row[18]),
                "sep_relevance": text(row[19]),
                "mitigation_ids": parse_refs(row[23]),
                "prop_rules": parse_refs(row[35]),
                "legacy": {
                    "text": legacy_text,
                    "scale": parse(Scale, lm.group("scale")),
                    "scope": parse(Scope, lm.group("scope")),
                    "reversibility": parse(Reversibility, lm.group("rev")),
                    "probability": parse(Probability, lm.group("prob")),
                    "reference_risk": parse(RiskLevel, lm.group("risk")),
                    "provisional_residual_text": text(row[30]),
                },
            }
        )
    return risks, read_calibration(ws)


def _row_starting(ws, prefix: str) -> int:
    for row in ws.iter_rows():
        v = text(row[0].value)
        if v and v.startswith(prefix):
            return row[0].row
    raise ExportError(f"{ws.title}: no row starting with {prefix!r}")


def read_calibration(ws) -> dict[str, Any]:
    def cells(r: int, n: int) -> list[Any]:
        return [ws.cell(r, c).value for c in range(1, n + 1)]

    rule_text = text(ws.cell(_row_starting(ws, "Impact severity = MAX"), 1).value)
    if "MAX(scale, scope, reversibility)" not in rule_text:
        raise ExportError(f"calibration rule changed: {rule_text!r}")
    m_row = _row_starting(ws, "Severity \\ Probability")
    probs = [parse(Probability, v) for v in cells(m_row, 5)[1:]]
    matrix: dict[Severity, dict[Probability, RiskLevel]] = {}
    for r in range(m_row + 1, m_row + 5):
        row = cells(r, 5)
        matrix[parse(Severity, row[0])] = {p: parse(RiskLevel, v) for p, v in zip(probs, row[1:])}
    d_row = _row_starting(ws, "Dimension")
    dims: dict[str, list[Any]] = {}
    for r, (name, enum) in enumerate(
        (("Scale", Scale), ("Scope", Scope), ("Reversibility", Reversibility), ("Probability", Probability)),
        start=d_row + 1,
    ):
        row = cells(r, 5)
        if text(row[0]) != name:
            raise ExportError(f"calibration: expected dimension {name} at row {r}, got {row[0]!r}")
        dims[name.lower()] = [parse(enum, v) for v in row[1:]]
    return {
        "rule_text": rule_text,
        "control_effect_text": text(ws.cell(_row_starting(ws, "CONTROL-EFFECT RULE"), 1).value),
        **dims,
        "matrix": matrix,
    }


def risk_id_of(cell: Any, where: str) -> str:
    m = re.match(r"^(RISK-\d{2})\b", text(cell) or "")
    if not m:
        raise ExportError(f"{where}: no risk id in {cell!r}")
    return m.group(1)


def read_mitigations(wb) -> list[dict[str, Any]]:
    ws = wb["05A_MITIGATION_PLAN"]
    hdr = find_header(ws, MITIGATION_HEADERS)
    out = []
    for row in data_rows(ws, hdr, 0, r"^MIT-\d{3}$"):
        mid = text(row[0])
        actual = text(row[23])
        before = text(row[21])
        out.append(
            {
                "id": mid,
                "order": len(out) + 1,
                "source_type": text(row[1]),
                "source_ids_text": text(row[2]),
                "risk_id": risk_id_of(row[3], mid),
                "source_rule": text(row[4]),
                "sep_ref": text(row[5]),
                "harm": text(row[6]),
                "rights": text(row[7]),
                "trigger": text(row[8]),
                "measure": text(row[9]),
                "hierarchy": parse(MitigationHierarchy, text(row[10])),
                "measure_type": text(row[11]),
                "responsible_role": text(row[12]),
                "owner": text(row[13]),
                "timing": lookup(TIMING, row[14], f"{mid} timing"),
                "timing_text": text(row[14]),
                "requirement": lookup(REQUIREMENT, row[15], f"{mid} requirement"),
                "decision_effect": lookup(DECISION_EFFECT, row[16], f"{mid} decision effect"),
                "decision_effect_text": text(row[16]),
                "evidence_of_implementation": text(row[17]),
                "verification_method": text(row[18]),
                "closure_criterion": text(row[19]),
                "status": parse(MitigationStatus, text(row[20])),
                "risk_before": parse(RiskLevel, before) if before else None,
                "expected_residual_text": text(row[22]),
                "actual_residual": parse(RiskLevel, actual) if actual else None,
                "adequacy_text": text(row[24]),
                "residual_acceptable_text": text(row[25]),
                "review_trigger": text(row[26]),
                "review_date": text(row[27]),
                "notes": text(row[28]),
                "expected_effect": text(row[29]),
                "reassessment_text": text(row[30]),
                "human_legal_review_text": text(row[31]),
                "unresolved_text": text(row[32]),
            }
        )
    return out


def read_sep(wb) -> dict[str, Any]:
    ws = wb["04A_SEP"]
    hdr = find_header(ws, SEP_HEADERS)
    fields = []
    for row in data_rows(ws, hdr, 0, r"^SEP\d{2}$"):
        fid = text(row[0])
        fields.append(
            {
                "id": fid,
                "order": len(fields) + 1,
                "field": text(row[1]),
                "question": text(row[2]),
                "response_model": text(row[3]),
                "activation": text(row[4]),
                "output": text(row[5]),
                "risk_link": text(row[6]),
                "rule_refs_text": text(row[7]),
                "rule_refs": parse_refs(row[7]),
                "mitigation_link": text(row[8]),
                "human_review": lookup(SEP_HUMAN_REVIEW, row[9], f"{fid} human review"),
                "human_review_text": text(row[9]),
                "status": text(row[10]),
            }
        )
    t_hdr = find_header(ws, ["Possible SEP trigger", "Default consequence", "Automatic?", "Notes"])
    triggers = []
    for row in ws.iter_rows(min_row=t_hdr + 1, values_only=True):
        if not text(row[0]):
            continue
        auto = lookup({"NO": False, "YES": True}, row[2], f"SEP trigger {row[0]!r}")
        triggers.append(
            {"trigger": text(row[0]), "default_consequence": text(row[1]), "automatic": auto, "notes": text(row[3])}
        )
    purpose = text(ws.cell(2, 1).value)
    if not (purpose and purpose.startswith("Purpose:")):
        raise ExportError("04A: purpose row not found")
    return {"purpose": purpose, "fields": fields, "triggers": triggers}


# ---------------------------------------------------------------------------
# assembly, validation, output
# ---------------------------------------------------------------------------


def registered_assumptions(path: Path) -> set[str]:
    if not path.exists():
        raise ExportError(f"{path} not found: the assumption register is required")
    return set(re.findall(r"\*\*(AS-\d{3})\*\*", path.read_text(encoding="utf-8")))


def build(workbook: Path) -> ConfigBundle:
    sha = hashlib.sha256(workbook.read_bytes()).hexdigest()
    meta = {
        "workbook_version": WORKBOOK_VERSION,
        "workbook_file": workbook.name,
        "workbook_sha256": sha,
        "generated_by": "tools/export_workbook.py",
    }
    wb = openpyxl.load_workbook(workbook, data_only=True, read_only=False)
    options = load_encoding("options.yaml")
    q_enc = load_encoding("questions.yaml")
    r_enc = load_encoding("rules.yaml")
    risks, calibration = read_risks(wb)
    sep = read_sep(wb)
    data = {
        "version": WORKBOOK_VERSION,
        "questions": {
            "meta": meta,
            "option_sets": [
                {"id": k, "source": "streamlit_app.py@896526b", "options": v, "assumptions": ["AS-001", "AS-008"]}
                for k, v in options.items()
            ],
            "questions": read_questions(wb, q_enc, options),
            "governance_fields": read_governance(wb),
        },
        "rules": {"meta": meta, "outcome_precedence": ["O5", "O4", "O3", "O2", "O1"], "rules": read_rules(wb, r_enc)},
        "risks": {"meta": meta, "calibration": calibration, "risks": risks},
        "mitigations": {"meta": meta, "mitigations": read_mitigations(wb)},
        "sep": {"meta": meta, **sep},
    }
    return ConfigBundle.model_validate(data)


def known_anomalies() -> list[dict[str, str]]:
    return load_encoding("known_anomalies.yaml") or []


def classify(findings: list[Finding], known: list[dict[str, str]]) -> list[tuple[str, Finding]]:
    out = []
    for f in findings:
        tag = f.level
        for k in known:
            if k["code"] == f.code and k["subject"] == f.subject and k.get("contains", "") in f.message:
                tag = f"KNOWN {k['ref']}"
                break
        out.append((tag, f))
    return out


HEADER = "# GENERATED by tools/export_workbook.py from {file} ({version}, sha256 {sha}).\n# Do not edit: change the workbook or tools/encodings/ and re-run the exporter.\n"


def render(bundle: ConfigBundle) -> dict[str, str]:
    files: dict[str, str] = {"VERSION": bundle.version + "\n"}
    meta = bundle.questions.meta
    head = HEADER.format(file=meta.workbook_file, version=bundle.version, sha=meta.workbook_sha256[:12])
    for name in ("questions", "rules", "risks", "mitigations", "sep"):
        payload = getattr(bundle, name).model_dump(mode="json", by_alias=True, exclude_none=True)
        body = yaml.safe_dump(payload, sort_keys=False, allow_unicode=True, width=110, default_flow_style=False)
        files[f"{name}.yaml"] = head + body
    return files


def encoding_report(bundle: ConfigBundle) -> str:
    lines = ["Rule        kind        assumptions"]
    for r in bundle.rules.rules:
        enc = r.encoding
        lines.append(f"{r.id:<11} {enc.kind.value:<11} {', '.join(enc.assumptions) or '-'}")
    kinds = Counter(r.encoding.kind.value for r in bundle.rules.rules)
    with_as = sum(1 for r in bundle.rules.rules if r.encoding.assumptions)
    lines.append(f"\nby kind: {dict(kinds)}; rules citing at least one AS-*: {with_as}/{len(bundle.rules.rules)}")
    return "\n".join(lines)


def short_errors(exc: ValidationError, limit: int = 15) -> list[str]:
    """One line per failing location; union alternatives collapsed to the shortest path."""
    seen: dict[str, str] = {}
    for err in exc.errors():
        loc = ".".join(str(p) for p in err["loc"])
        root = re.split(r"\.(?:when|gap|visibility)\.", loc)[0]
        if root not in seen or len(loc) < len(seen[root]):
            seen[root] = f"{loc}: {err['msg']} (input {str(err.get('input'))[:80]})"
    return list(seen.values())[:limit]


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    ap.add_argument("--anomalies", type=Path, default=DEFAULT_ANOMALIES)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--check", action="store_true", help="validate and compare with --out, write nothing")
    ap.add_argument("--encoding-report", action="store_true", help="print rule encodings and their assumptions")
    args = ap.parse_args(argv)

    try:
        bundle = build(args.workbook)
        registered = registered_assumptions(args.anomalies)
    except ValidationError as exc:
        print(f"[ERROR] export failed: {exc.error_count()} schema errors")
        for line in short_errors(exc):
            print(f"  {line}")
        return 1
    except (ExportError, yaml.YAMLError, ValueError) as exc:
        print(f"[ERROR] export failed: {exc}")
        return 1

    findings = classify(validate_bundle(bundle, registered), known_anomalies())
    order = {"ERROR": 0, "WARN": 2, "INFO": 3}
    findings.sort(key=lambda t: (order.get(t[0], 1), t[1].code, t[1].subject))
    for tag, f in findings:
        print(f"[{tag}] {f.code} {f.subject}: {f.message}")
    tags = Counter(t.split()[0] for t, _ in findings)
    q = bundle.questions.questions
    print(
        f"\n{bundle.version}: {len(q)} questions ({Counter(x.phase.value for x in q)}), "
        f"{len(bundle.rules.rules)} rules, {len(bundle.risks.risks)} risks, "
        f"{len(bundle.mitigations.mitigations)} mitigations, {len(bundle.sep.fields)} SEP fields"
    )
    print(f"findings: {dict(tags)}")
    if args.encoding_report:
        print("\n" + encoding_report(bundle))
    if tags.get("ERROR"):
        print("\nnew ERROR findings: config/ not written")
        return 1

    files = render(bundle)
    if args.check:
        stale = [n for n, body in files.items() if not (args.out / n).exists() or (args.out / n).read_text(encoding="utf-8") != body]
        print("config/ is up to date" if not stale else f"config/ is stale: {stale}")
        return 1 if stale else 0
    args.out.mkdir(parents=True, exist_ok=True)
    for name, body in files.items():
        (args.out / name).write_text(body, encoding="utf-8", newline="\n")
    print(f"wrote {', '.join(files)} to {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
